"""Tests for the NSZ viewer: the helpers on synthetic data, and the parser and the application against the
real Horiba SZ-100 sample in 'example data/' (an .nsz file plus the CSV exported by the instrument software for the
same measurement; gitignored, real data). Tests that need the sample are skipped if it is missing.

    py -m unittest discover -s tests -v
"""
import csv
import importlib.util
import os
import pickle
import struct
import tempfile
import tkinter as tk
import unittest
from unittest import mock

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("nsz", os.path.join(HERE, "..", "nsz_viewer.pyw"))
nsz = importlib.util.module_from_spec(spec)
spec.loader.exec_module(nsz)

DATA = os.path.join(HERE, "..", "example data")
SAMPLE = os.path.join(DATA, "lipo_bup_MLV_0minson_scattering_0729.nsz")
SAMPLE_CSV = os.path.join(DATA, "lipo_bup_MLV_0minson_scattering_0731.csv")
HAS_SAMPLE = os.path.isfile(SAMPLE) and os.path.isfile(SAMPLE_CSV)


def read_software_csv(path):
    """The instrument software's export: (item, value) rows, then three numeric tables separated by blank lines:
    the size distribution, the residual g1 and the correlation g1 with its fit."""
    with open(path, newline="", encoding="latin-1") as f:
        rows = list(csv.reader(f))
    items = {r[0]: r[1] for r in rows if len(r) >= 2 and r[0] and not r[0][0].isdigit()}

    def table_after(header_start):
        start = next(i for i, r in enumerate(rows) if r and ",".join(r).startswith(header_start)) + 1
        out = []
        for r in rows[start:]:
            if not r:
                break
            out.append([float(x) for x in r])
        return np.array(out)

    return items, table_after("Diameter (nm)"), table_after("Delay Time(usec),Residual"),         table_after("Delay Time(usec),Correlation")


def synthetic_sample(name="synthetic"):
    """A sample dict shaped like load_nsz() output, without any file."""
    tau = np.logspace(0, 6, 60)
    acf = np.exp(-tau / 2e4) * 0.9
    sz = np.logspace(-0.5, 3.8, 84)
    intens = np.exp(-0.5 * ((np.log10(sz) - 2.5) / 0.15) ** 2)
    intens = intens / intens.sum() * 100
    return {
        'path': name + ".nsz", 'name': name, 'tau': tau, 'acf': acf,
        'fit': acf[6:] * 0.99, 'resid': acf[6:] * 0.01, 'fit_offset': 6,
        'sz': sz, 'intens': intens, 'cumul': np.cumsum(intens),
        'params': {'Sample_Name': name, 'Date': '2026-01-02', 'Time': '10:11:12', 'D50_nm': 316.0,
                   'D10_nm': 250.0, 'D90_nm': 400.0, 'Peak1_nm': 316.0, 'PDI': 0.1, 'Peak2_nm': None},
    }


class TestHelpers(unittest.TestCase):
    def test_f32(self):
        data = struct.pack("<4f", 1.0, 2.5, -3.0, 4.0) + b"\x00\x01"       # trailing bytes are ignored
        self.assertEqual(nsz.f32(data), [1.0, 2.5, -3.0, 4.0])

    def test_parse_props(self):
        data = b"junk\x00" + b"MeasAngle\x00" + struct.pack("<f", 173.0) + b"\x00" \
               + b"MeasHolderTemp\x00" + struct.pack("<f", 24.5) + b"\x00" + b"tail" * 4
        props = nsz.parse_props(data)
        self.assertAlmostEqual(props["MeasAngle"], 173.0)
        self.assertAlmostEqual(props["MeasHolderTemp"], 24.5)

    def test_number_distribution(self):
        sz = np.array([0.0, 10.0, 100.0, 1000.0])
        intens = np.array([5.0, 10.0, 10.0, 10.0])
        num, cumul = nsz.NSZViewer._to_number_dist(sz, intens)
        self.assertAlmostEqual(num.sum(), 100.0)
        self.assertEqual(num[0], 0.0)                         # zero diameter contributes nothing
        self.assertGreater(num[1], num[2])                    # N is proportional to I / d^6: small particles dominate
        self.assertAlmostEqual(cumul[-1], 100.0)
        self.assertTrue(np.all(np.diff(cumul) >= 0))

    def test_missing_olefile_message(self):
        with mock.patch.object(nsz, "OLEFILE_AVAILABLE", False):
            with self.assertRaises(ImportError):
                nsz.load_nsz("whatever.nsz")

    def test_sample_colors_are_a_cycle(self):
        self.assertGreaterEqual(len(nsz.sample_colors()), 2)

    def test_excel_export(self):
        if not nsz.XLSX_AVAILABLE:
            self.skipTest("openpyxl not installed")
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, "x.xlsx")
            nsz.NSZViewer._write_xlsx(out, synthetic_sample())
            wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ["ACF", "Intensity distribution", "Number distribution", "Parameters"])
        self.assertEqual([c.value for c in wb["ACF"][1]], ["tau (µs)", "g2-1 exp", "g2-1 fit", "Residual"])
        self.assertEqual(wb["Intensity distribution"].max_row, 85)            # header + 84 bins
        labels = [r[0].value for r in wb["Parameters"].iter_rows(min_row=2)]
        self.assertIn("Sample name", labels)
        self.assertIn("D50 (nm)", labels)


@unittest.skipUnless(HAS_SAMPLE, "real sample (.nsz + instrument CSV) not present")
class TestRealSample(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = nsz.load_nsz(SAMPLE)
        cls.items, cls.table, cls.resid_csv, cls.corr_csv = read_software_csv(SAMPLE_CSV)

    def test_results_match_the_software(self):
        p = self.d["params"]
        self.assertAlmostEqual(p["Mean1_nm"], float(self.items["Total Mean"]), delta=0.06)
        self.assertAlmostEqual(p["SD1_nm"], float(self.items["Total S. D."]), delta=0.06)
        self.assertAlmostEqual(p["Mode_nm"], float(self.items["Total Mode"]), delta=0.06)
        self.assertAlmostEqual(p["ZAvg_nm"], float(self.items["Z-Average"]), delta=0.06)
        self.assertAlmostEqual(p["PDI"], float(self.items["PI"]), delta=0.0006)

    def test_measurement_parameters(self):
        p = self.d["params"]
        self.assertEqual(p["Angle_deg"], float(self.items["Scattering Angle"]))
        self.assertAlmostEqual(p["Temp_C"], float(self.items["Temperature of the Holder"]), delta=0.06)
        self.assertAlmostEqual(p["Visc_mPas"], float(self.items["Dispersion Medium Viscosity"]), delta=0.001)
        self.assertEqual(p["Sample_Name"], self.items["Sample Name"])

    def test_size_distribution_matches_the_software_table(self):
        self.assertEqual(len(self.d["sz"]), len(self.table))
        np.testing.assert_allclose(self.d["sz"], self.table[:, 0], atol=0.006)     # the CSV rounds to 2 decimals
        np.testing.assert_allclose(self.d["intens"], self.table[:, 1], atol=0.0006)
        np.testing.assert_allclose(self.d["cumul"], self.table[:, 2], atol=0.0006)

    def test_correlation_matches_the_software_g1_squared(self):
        # The software exports g1; the file stores g1 squared. CSV row k is channel 7 + k (the 6 channels excluded from
        # the fit plus one). The viewer drops the trailing baseline channels, so only the first rows can be compared;
        # the CSV also floors g1 at 0.01, so compare only the values above the floor.
        acf = self.d["acf"][7:]
        g1_csv = self.corr_csv[:len(acf), 1]
        self.assertGreaterEqual(len(acf), 25)
        ok = g1_csv > 0.011
        np.testing.assert_allclose(np.sqrt(np.clip(acf[ok], 0, None)), g1_csv[ok], atol=2e-5)

    def test_fit_matches_the_software_g1_squared(self):
        fit_csv = self.corr_csv[:, 2]
        n = int((fit_csv > 0).sum())                          # the fit stops where the software fit is 0
        fit = self.d["fit"][1:1 + n]
        np.testing.assert_allclose(np.sqrt(fit), fit_csv[:n], atol=2e-5)

    def test_residuals_follow_the_software_pattern(self):
        # Residuals are of g1 in the CSV and of g1 squared in the file, so they are not equal: check that they are
        # clearly correlated at the expected alignment (CSV row k = residual k + 1) and worse at the neighbours.
        theirs = self.resid_csv[:, 1]

        def corr(offset):
            n = min(len(theirs), len(self.d["resid"]) - offset)
            return np.corrcoef(self.d["resid"][offset:offset + n], theirs[:n])[0, 1]

        self.assertGreater(corr(1), 0.75)
        self.assertGreater(corr(1), max(corr(0), corr(2)))

    def test_correlation_function_shape(self):
        d = self.d
        self.assertEqual(len(d["tau"]), len(d["acf"]))
        self.assertTrue(np.all(np.diff(d["tau"]) > 0))        # the repeated baseline channels were trimmed
        self.assertEqual(d["fit_offset"], 6)
        self.assertEqual(len(d["fit"]), len(d["resid"]))
        self.assertLessEqual(len(d["fit"]) + d["fit_offset"], len(d["tau"]))


class TestApplication(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            cls.root = nsz.TkinterDnD.Tk() if nsz.DND_AVAILABLE else tk.Tk()
        except tk.TclError as e:
            raise unittest.SkipTest("Tk not available: %s" % e)
        cls.root.withdraw()

    @classmethod
    def tearDownClass(cls):
        cls.root.destroy()

    def setUp(self):
        self.app = nsz.NSZViewer(self.root)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def tearDown(self):
        for w in self.root.winfo_children():
            if isinstance(w, tk.Toplevel):
                w.destroy()

    def add_synthetic(self, *names):
        for n in names:
            d = synthetic_sample(n)
            self.app.samples[d["path"]] = d
            self.app.listbox.insert(tk.END, n + ".nsz")
        self.app.listbox.selection_clear(0, tk.END)
        for i in range(self.app.listbox.size()):
            self.app.listbox.selection_set(i)
        self.app._on_select()
        self.root.update()

    def test_empty_state(self):
        self.assertEqual(self.app.selected, [])
        self.assertIsNone(self.app._current_figure())

    def test_selecting_samples_draws_the_plots(self):
        self.add_synthetic("a", "b")
        self.assertEqual(len(self.app.selected), 2)
        self.assertEqual(len(self.app.fig_acf.axes[0].lines), 2 * 2 + 1)      # 2 samples: data + fit, plus the zero line
        self.assertEqual(len(self.app.fig_size.axes), 2)                       # distribution + cumulative

    def test_residuals_and_number_representation(self):
        self.add_synthetic("a")
        self.app.show_resid.set(True)
        self.app._plot_acf()
        self.assertEqual(len(self.app.fig_acf.axes), 2)
        self.app.size_repr.set("number")
        self.app.show_cumul.set(False)
        self.app._plot_size()
        self.assertEqual(len(self.app.fig_size.axes), 1)
        self.assertEqual(self.app.fig_size.axes[0].get_ylabel(), "Number (%)")

    def test_parameters_table(self):
        self.add_synthetic("a", "b")
        rows = [self.app.tree.item(i, "values") for i in self.app.tree.get_children()]
        by_label = {r[0]: r for r in rows}
        self.assertEqual(by_label["D50 (nm)"][1:3], ("316.0", "316.0"))
        self.assertEqual(by_label["Peak 2 (nm)"][1], "—")                      # missing value
        self.assertEqual(self.app.tree.heading("Sample 1")["text"], "a")

    def test_remove_selected(self):
        self.add_synthetic("a", "b")
        self.app.listbox.selection_clear(0, tk.END)
        self.app.listbox.selection_set(0)
        self.app._remove_selected()
        self.assertEqual(list(self.app.samples), ["b.nsz"])
        self.assertEqual(self.app.listbox.size(), 1)

    def test_figure_follows_the_active_tab(self):
        self.add_synthetic("a")
        self.app.notebook.select(self.app.tab_acf)
        self.assertIs(self.app._current_figure(), self.app.fig_acf)
        self.app.notebook.select(self.app.tab_size)
        self.assertIs(self.app._current_figure(), self.app.fig_size)
        self.app.notebook.select(self.app.tab_params)
        self.assertIsNone(self.app._current_figure())

    def test_figure_copy_and_save(self):
        self.add_synthetic("a")
        self.app.notebook.select(self.app.tab_size)
        copy = self.app._copy_figure(self.app.fig_size)
        self.assertEqual(len(copy.axes), 2)
        if nsz.origin_style is not None:
            w, h = copy.get_size_inches()
            self.assertAlmostEqual(w / h, 16 / 9, places=2)                    # two panels -> 'double' preset
        out = os.path.join(self.tmp.name, "f.fig.pickle")
        with mock.patch.object(nsz.filedialog, "asksaveasfilename", return_value=out), \
                mock.patch.object(nsz.messagebox, "showinfo"), \
                mock.patch.object(nsz.messagebox, "showerror") as err:
            self.app.save_figure_pickle()
        self.assertFalse(err.called)
        with open(out, "rb") as f:
            self.assertEqual(len(pickle.load(f).axes), 2)
        png = os.path.join(self.tmp.name, "f.png")
        with mock.patch.object(nsz.filedialog, "asksaveasfilename", return_value=png), \
                mock.patch.object(nsz.messagebox, "showinfo"):
            self.app.save_figure_image()
        self.assertGreater(os.path.getsize(png), 5000)

    def test_figure_actions_without_a_plot_warn(self):
        with mock.patch.object(nsz.messagebox, "showwarning") as warn:
            self.app.save_figure_image()
            self.app.save_figure_pickle()
            self.app.open_figure_editor()
        self.assertEqual(warn.call_count, 3)

    def test_export_single_and_several(self):
        if not nsz.XLSX_AVAILABLE:
            self.skipTest("openpyxl not installed")
        self.add_synthetic("a")
        out = os.path.join(self.tmp.name, "one.xlsx")
        with mock.patch.object(nsz.filedialog, "asksaveasfilename", return_value=out), \
                mock.patch.object(nsz.messagebox, "showinfo"):
            self.app._export_xlsx()
        self.assertTrue(os.path.isfile(out))
        self.add_synthetic("b")
        with mock.patch.object(nsz.filedialog, "askdirectory", return_value=self.tmp.name), \
                mock.patch.object(nsz.messagebox, "showinfo"):
            self.app._export_xlsx()
        self.assertTrue(os.path.isfile(os.path.join(self.tmp.name, "b.xlsx")))

    def test_loading_a_bad_file_reports_the_error(self):
        bad = os.path.join(self.tmp.name, "bad.nsz")
        with open(bad, "wb") as f:
            f.write(b"not an OLE file")
        with mock.patch.object(nsz.messagebox, "showerror") as err:
            self.app._load_files([bad])
        self.assertTrue(err.called)
        self.assertEqual(self.app.samples, {})

    @unittest.skipUnless(HAS_SAMPLE, "real sample not present")
    def test_real_file_through_the_application(self):
        self.app._load_files([SAMPLE])
        self.root.update()
        self.assertEqual(self.app.listbox.size(), 1)
        self.assertEqual(len(self.app.selected), 1)
        rows = {self.app.tree.item(i, "values")[0]: self.app.tree.item(i, "values")
                for i in self.app.tree.get_children()}
        self.assertEqual(rows["Z-Average (nm)"][1], "1117.7646")
        self.assertEqual(rows["PDI"][1], "0.6013")
        self.app._load_files([SAMPLE])                                          # same file again: ignored
        self.assertEqual(self.app.listbox.size(), 1)


if __name__ == "__main__":
    unittest.main()
